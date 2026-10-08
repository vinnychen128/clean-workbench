# SPDX-License-Identifier: Apache-2.0
"""L1 单元测试：解析引擎。

覆盖：类型白名单 / 50MB 上限 / 空文件 / CSV 编码探测(UTF-8·GBK·GB18030) /
分隔符探测 / 长数字保持文本(长数字⚠) / Excel(xlsx/xls) / 扫描版 PDF 拒绝 /
同一性哈希。
"""
import base64

import pytest

from app.parsing.parser import (
    ALLOWED_TYPES,
    MAX_FILE_SIZE,
    ParseError,
    compute_file_hash,
    parse_csv,
    parse_excel,
    parse_pdf,
    parse_stream,
)

CSV_UTF8 = (
    "id,名称,金额,日期,备注\n"
    "A001,张三,￥1,200,2024/01/01,正常\n"
    "A002,李四,3000,2024-02-01, 有空格 \n"
)


def b64(s: str) -> str:
    return base64.b64encode(s.encode("utf-8")).hex()  # placeholder


def _b64b(s: str) -> str:
    return base64.b64encode(s.encode("utf-8")).decode()


# --- 类型白名单 / 大小 / 空文件 ------------------------------------------
def test_allowed_types():
    assert ALLOWED_TYPES == {"csv", "xlsx", "xls", "pdf"}


def test_unsupported_type():
    with pytest.raises(ParseError) as ei:
        parse_stream("a.exe", "exe", 100, "", _b64b("x"), "t")
    assert ei.value.code == "UNSUPPORTED_TYPE"


def test_too_large():
    with pytest.raises(ParseError) as ei:
        parse_stream("a.csv", "csv", MAX_FILE_SIZE + 1, "", _b64b("x"), "t")
    assert ei.value.code == "FILE_TOO_LARGE"


def test_empty_file():
    with pytest.raises(ParseError) as ei:
        parse_stream("a.csv", "csv", 0, "", _b64b(""), "t")
    assert ei.value.code == "EMPTY_FILE"


# --- CSV 解析 -------------------------------------------------------------
def test_parse_csv_basic():
    pr = parse_csv(CSV_UTF8.encode("utf-8"), "ref://t")
    assert pr.columns == ["id", "名称", "金额", "日期", "备注"]
    assert pr.row_count == 2
    assert pr.col_count == 5
    assert pr.parsing_report["delimiter"] == ","


def test_parse_csv_gbk():
    raw = CSV_UTF8.encode("gbk")
    pr = parse_csv(raw, "ref://t")
    assert pr.columns[1] == "名称"
    assert pr.rows[0][1] == "张三"


def test_parse_csv_gb18030():
    raw = "id,name\n1,中文𠀀\n".encode("gb18030")
    pr = parse_csv(raw, "ref://t")
    assert pr.rows[0][1] == "中文𠀀"


def test_parse_csv_tab_delimiter():
    pr = parse_csv(b"a\tb\n1\t2\n", "ref://t")
    assert pr.columns == ["a", "b"]
    assert pr.rows == [["1", "2"]]


def test_parse_csv_semicolon_delimiter():
    pr = parse_csv(b"a;b\n1;2\n", "ref://t")
    assert pr.columns == ["a", "b"]


def test_long_number_kept_as_text():
    raw = "id,卡号\n1,12345678901234567890\n".encode("utf-8")
    pr = parse_csv(raw, "ref://t")
    assert pr.rows[0][1] == "12345678901234567890"  # 原样文本，不丢精度


def test_excel_long_number_warns(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["卡号"])
    ws.append([123456789012345678])
    p = tmp_path / "t.xlsx"
    wb.save(p)
    pr = parse_excel(p.read_bytes(), "xlsx", "ref://t")
    # Excel 双精度本身会丢精度：读回为科学计数 float，被文本化并加长数字⚠
    assert "长数字⚠" in pr.parsing_report["warnings"]
    assert isinstance(pr.rows[0][0], str)


def test_parse_csv_blank_row_removed():
    raw = "a,b\n1,2\n\n , \n3,4\n".encode("utf-8")
    pr = parse_csv(raw, "ref://t")
    assert pr.row_count == 2


# --- Excel 解析 -----------------------------------------------------------
def test_parse_excel_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "金额"])
    ws.append([1, 100])
    ws.append([2, None])
    p = tmp_path / "t.xlsx"
    wb.save(p)
    pr = parse_excel(p.read_bytes(), "xlsx", "ref://t")
    assert pr.columns == ["id", "金额"]
    assert pr.row_count == 2
    assert pr.parsing_report["selected_sheet"] == "Sheet1"


def test_parse_excel_xls(tmp_path):
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("S1")
    ws.write(0, 0, "id")
    ws.write(0, 1, "金额")
    ws.write(1, 0, 1)
    ws.write(1, 1, 200)
    p = tmp_path / "t.xls"
    wb.save(str(p))
    pr = parse_excel(p.read_bytes(), "xls", "ref://t")
    assert pr.columns == ["id", "金额"]
    assert pr.row_count == 1


# --- PDF 解析（不产出假数据） ---------------------------------------------
def test_parse_pdf_no_fabrication():
    raw = b"%PDF-1.4 minimal"
    try:
        pr = parse_pdf(raw, "ref://t")
        # 文本 PDF：只登记引用与提示，不产出假数据
        assert pr.row_count == 0
        assert any("不接收原件" in w for w in pr.parsing_report["warnings"])
    except ParseError as exc:
        # 扫描件 / 无解析库：明确拒绝，同样不产出假数据
        assert exc.code in ("SCANNED_PDF_UNSUPPORTED", "PARSE_FAILED")


# --- 哈希与 parse_stream ---------------------------------------------------
def test_compute_file_hash_deterministic():
    assert compute_file_hash(b"abc") == compute_file_hash(b"abc")
    assert compute_file_hash(b"abc") != compute_file_hash(b"abd")


def test_parse_stream_csv():
    raw = CSV_UTF8.encode("utf-8")
    pr = parse_stream("orders.csv", "csv", len(raw), "", _b64b(CSV_UTF8), "t1")
    assert pr.row_count == 2
    assert pr.data_ref == "ref://parse/t1"
